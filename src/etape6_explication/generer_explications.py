"""
generer_explications.py (Etape 6c - orchestrateur)

Pour chaque KPI du classement final (resultats/top_kpi.csv) :
1. Retrouve les colonnes brutes qui le composent et leurs VRAIES definitions
   metier (dictionnaire_enrichi.json, issu de l'equipe metier - etapes 1-2).
2. Calcule un effet XAI objectif par quartile (xai_effets.py, aucun modele
   de langage).
3. Demande au modele de langage de traduire tout ca en explication metier
   (prompt_explication.py), en lui interdisant d'inventer un sens qui ne
   decoule pas des definitions fournies, et en lui imposant de signaler
   les KPI sans pouvoir predictif demontre (signal_predictif_significatif == False).

INPUT  : resultats/top_kpi.csv (etape 5) + dictionnaire_enrichi.json (etape 2)
         + le jeu de donnees actif (config.DATA_PATH)
OUTPUT : resultats/top_kpi_explique.xlsx
"""
import json
import re
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT_DIR))
sys.path.insert(0, str(ROOT_DIR / "src" / "etape4_candidats"))
sys.path.insert(0, str(ROOT_DIR / "src" / "etape5_notation"))

from config import (  # noqa: E402
    TARGET_COL, MODEL_NAME, OLLAMA_BASE_URL, DATA_GENERATED_DIR,
    TOP_KPI_PATH, TOP_KPI_EXPLIQUE_PATH,
)
from charger_classification import load_working_lists  # noqa: E402
from generation_candidats import load_real_data, generate_all_candidates  # noqa: E402
from generation_candidats_guides import load_propositions, generate_guided_candidates  # noqa: E402
from notation_kpi import human_readable_name  # noqa: E402

from xai_effets import compute_quartile_effect  # noqa: E402
from prompt_explication import SYSTEM_PROMPT, build_kpi_prompt  # noqa: E402

DICT_ENRICHI_PATH = DATA_GENERATED_DIR / "dictionnaire_enrichi.json"
OUTPUT_PATH = TOP_KPI_EXPLIQUE_PATH

NUM_CTX = 8192
MAX_RETRIES = 2
TEMPERATURE = 0.2  # legerement au-dessus de la classification (0.1) : redaction plus naturelle, toujours peu creative


def load_definitions(dict_path: Path = DICT_ENRICHI_PATH) -> dict:
    with open(dict_path, encoding="utf-8") as f:
        records = json.load(f)
    return {r["nom"]: r["definition"] for r in records}


def extract_source_columns(candidate_key: str) -> list[str]:
    """Meme decoupage que human_readable_name (notation_kpi.py), mais
    retourne les noms de colonnes brutes plutot qu'une phrase."""
    m = re.match(r"^ratio__(.+)__sur__(.+)$", candidate_key)
    if m:
        return list(m.groups())
    m = re.match(r"^moyenne__(.+)__par__(.+)$", candidate_key)
    if m:
        return list(m.groups())
    m = re.match(r"^nombre_observations__par__(.+)$", candidate_key)
    if m:
        return [m.group(1)]
    m = re.match(r"^delai__(.+)__moins__(.+)$", candidate_key)
    if m:
        return list(m.groups())
    return []


def base_column_name(col: str) -> str:
    """Retire les suffixes __annee/__mois des dimensions de date derivees
    (voir derive_date_parts, generation_candidats.py) pour retrouver le nom
    de la colonne source dans le dictionnaire."""
    for suffix in ("__annee", "__mois"):
        if col.endswith(suffix):
            return col[: -len(suffix)]
    return col


def infos_candidat_metier(proposition: dict) -> tuple[list[str], str]:
    """Equivalent de extract_source_columns + human_readable_name, mais pour
    un candidat issu du raisonnement metier (etape 3b) : la formule vient
    directement de la proposition structuree, pas d'un decoupage du nom
    technique par regex."""
    operation = proposition.get("operation")

    if operation == "ratio":
        num, denom = proposition["numerateur"], proposition["denominateur"]
        colonnes = num + denom
        formule = f"({' + '.join(num)}) / ({' + '.join(denom)})"
    elif operation == "brute":
        colonnes = [proposition["colonne"]]
        formule = f"Colonne existante : {proposition['colonne']}"
    elif operation == "complement":
        colonnes = [proposition["colonne"]]
        formule = f"1 - {proposition['colonne']}"
    elif operation == "somme_ratios":
        colonnes = []
        parts = []
        for r in proposition["ratios"]:
            colonnes += r["numerateur"] + r["denominateur"]
            parts.append(f"({' + '.join(r['numerateur'])}) / ({' + '.join(r['denominateur'])})")
        formule = " + ".join(parts)
    else:
        colonnes, formule = [], "(formule inconnue)"

    return colonnes, formule


def call_ollama(prompt: str, model: str = MODEL_NAME) -> str:
    import ollama
    client = ollama.Client(host=OLLAMA_BASE_URL)
    response = client.chat(
        model=model,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": prompt},
        ],
        format="json",
        options={"temperature": TEMPERATURE, "num_ctx": NUM_CTX},
    )
    return response["message"]["content"]


def extract_json(raw: str) -> dict:
    raw = raw.strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip())
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        pass
    match = re.search(r"\{.*\}", raw, re.DOTALL)
    if match:
        return json.loads(match.group(0))
    raise json.JSONDecodeError("Aucun bloc JSON trouve", raw, 0)


def explain_kpi_avec_formule(nom_lisible: str, formule: str, definitions: dict, scores: dict, xai: dict) -> str:
    prompt = build_kpi_prompt(nom_lisible, formule, definitions, scores, xai)

    last_raw = None
    for attempt in range(1, MAX_RETRIES + 1):
        raw = call_ollama(prompt)
        last_raw = raw
        try:
            parsed = extract_json(raw)
            explication = parsed.get("explication")
            if explication:
                return explication.strip()
        except json.JSONDecodeError:
            pass
        print(f"    [tentative {attempt}/{MAX_RETRIES}] reponse invalide, reponse brute : {last_raw[:200]!r}")
        time.sleep(1)

    return "(echec de generation - reponse du modele invalide apres plusieurs tentatives)"


def build_excel(rows: list[dict]) -> None:
    df_out = pd.DataFrame(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "KPI expliques"

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    normal_font = Font(name="Arial", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, col_name in enumerate(df_out.columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font, cell.fill, cell.alignment = header_font, header_fill, wrap

    for row_idx, row in enumerate(df_out.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font, cell.alignment = normal_font, wrap

    widths = [45, 18, 12, 12, 14, 12, 12, 16, 55]
    for i, w in enumerate(widths, start=1):
        if i <= len(df_out.columns):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)


def main():
    top_kpis = pd.read_csv(TOP_KPI_PATH)
    definitions = load_definitions()

    print("Regeneration des candidats (necessaire pour l'effet XAI)...")
    df = load_real_data()
    working_lists = load_working_lists(df=df)
    candidates_exhaustifs, _ = generate_all_candidates(df, working_lists)

    propositions = load_propositions()
    candidates_metier, infos_metier = generate_guided_candidates(df, propositions)

    candidates = pd.concat([candidates_exhaustifs, candidates_metier], axis=1)
    target = df[TARGET_COL]

    rows = []
    for i, kpi in top_kpis.iterrows():
        candidate_key = kpi["candidat_technique"]
        origine = kpi.get("origine", "generation_exhaustive")
        print(f"\n[{i + 1}/{len(top_kpis)}] {kpi['nom_lisible']} ({origine})")

        if candidate_key not in candidates.columns:
            print("  ATTENTION : candidat introuvable dans les donnees regenerees, ignore.")
            continue

        if origine == "raisonnement_metier" and candidate_key in infos_metier:
            source_cols_bruts, formule = infos_candidat_metier(infos_metier[candidate_key])
        else:
            source_cols_bruts, formule = extract_source_columns(candidate_key), human_readable_name(candidate_key)

        source_cols = sorted({base_column_name(c) for c in source_cols_bruts})
        kpi_definitions = {c: definitions.get(c, "(definition non disponible)") for c in source_cols}

        xai = compute_quartile_effect(candidates[candidate_key], target)

        scores = {
            "score_composite": round(float(kpi["score_composite"]), 4),
            "score_predictif": round(float(kpi["score_predictif"]), 4),
            "signal_predictif_significatif": bool(kpi["signal_predictif_significatif"]),
            "score_information": round(float(kpi["score_information"]), 4),
            "score_stabilite": round(float(kpi["score_stabilite"]), 4),
            "completude": round(float(kpi["completude"]), 4),
            "origine": origine,
        }

        explication = explain_kpi_avec_formule(kpi["nom_lisible"], formule, kpi_definitions, scores, xai)
        print(f"  -> {explication[:120]}{'...' if len(explication) > 120 else ''}")

        rows.append({
            "KPI": kpi["nom_lisible"],
            "origine": origine,
            "score_composite": scores["score_composite"],
            "score_predictif": scores["score_predictif"],
            "signal_predictif_significatif": scores["signal_predictif_significatif"],
            "score_stabilite": scores["score_stabilite"],
            "completude": scores["completude"],
            "effet_xai": (
                f"Bas={xai['mediane_cible_quartile_bas']} / Haut={xai['mediane_cible_quartile_haut']} "
                f"({xai['direction']})" if xai.get("calculable") else "non calculable"
            ),
            "explication_metier": explication,
        })

    build_excel(rows)
    print(f"\n{len(rows)} KPI expliques -> {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
